#!/usr/bin/env python3
"""Build editable PPTX decks for every "Wait, I Qualify?!" graphic, for import
into Canva (PPTX import keeps text boxes, shapes and images editable).

    uv run --with python-pptx show/assets/templates/build_templates.py

Writes to show/assets/templates/out/:
  wiq-series-template-kit.pptx   one page per template (intro, lower thirds,
                                 program cards, URL strap, book card, Find Help,
                                 outro end card, chart frame)
  wiq-epN-<slug>.pptx            per-episode deck, cards pre-filled from the
                                 script cues in episodes.py
  wiq-youtube-thumbnails.pptx    one 16:9 page per episode
  techempower-youtube-banner.pptx / techempower-youtube-avatar.pptx

Geometry is in 1920x1080 pixel space (144 px per inch on a 13.333x7.5in slide);
colours are the Ep1 overlay palette (show/assets/generate-ep1-overlays.py) and
the site design tokens (styles/global.css). Fonts are named only -- Canva maps
Fraunces and DM Sans to its own library on import.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).parent))
from episodes import EPISODES, THUMB_ONLY  # noqa: E402

HERE = Path(__file__).parent
OUT = HERE / "out"
SUN = HERE.parent / "logo" / "techempower-sun.png"

DISPLAY = "Fraunces"
SANS = "DM Sans"

# --- palette -----------------------------------------------------------------
CREAM = RGBColor(0xFD, 0xFB, 0xF8)
CARD = RGBColor(0xFD, 0xF6, 0xEC)
PLACEHOLDER = RGBColor(0xED, 0xE7, 0xDD)
BARK = RGBColor(0x2C, 0x1F, 0x1A)
BROWN = RGBColor(0x4A, 0x38, 0x28)
DARK = RGBColor(0x3A, 0x2A, 0x1F)
INK = RGBColor(0x5A, 0x4E, 0x42)
MUTED = RGBColor(0x8A, 0x7A, 0x6C)
RULE = RGBColor(0xD5, 0xCD, 0xC0)
TEAL_BOX = RGBColor(0x1D, 0x5C, 0x5C)
TEAL = RGBColor(0x2A, 0x7F, 0x7F)
TEAL_SITE = RGBColor(0x0F, 0x76, 0x6E)
AMBER = RGBColor(0xD4, 0xA0, 0x3C)
AMBER_DK = RGBColor(0xB4, 0x53, 0x09)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFFWHITE = RGBColor(0xE2, 0xE5, 0xE2)
PALE = RGBColor(0xFB, 0xF6, 0xEE)

W, H = 1920, 1080


def E(px: float) -> Emu:
    """1920px == 13.333in -> 6350 EMU per pixel."""
    return Emu(int(round(px * 6350)))


# --- primitives --------------------------------------------------------------
def new_prs(square: bool = False) -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(7.5) if square else Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def page(prs, bg: RGBColor | None = None, name: str | None = None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    if bg is not None:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = bg
    if name:
        slide.name = name
    return slide


def rect(slide, x, y, w, h, fill, *, name=None, radius=None, edge=None,
         edge_w=2, dashed=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(kind, E(x), E(y), E(w), E(h))
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if edge is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = edge
        shp.line.width = E(edge_w)
        if dashed:
            shp.line.dash_style = MSO_LINE.DASH
    if radius:
        shp.adjustments[0] = min(0.5, radius / min(w, h))
    if name:
        shp.name = name
    return shp


def dot(slide, x, y, d, fill, name=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, E(x), E(y), E(d), E(d))
    shp.shadow.inherit = False
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    if name:
        shp.name = name
    return shp


def text(slide, x, y, w, h, lines, *, size, color, font=SANS, bold=False,
         italic=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name=None,
         spacing=None):
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.line_spacing = spacing
        r = p.add_run()
        r.text = line
        f = r.font
        f.name = font
        f.size = Pt(size)
        f.bold = bold
        f.italic = italic
        f.color.rgb = color
    if name:
        tb.name = name
    return tb


def picture(slide, path, x, y, w, h, name=None):
    pic = slide.shapes.add_picture(str(path), E(x), E(y), E(w), E(h))
    if name:
        pic.name = name
    return pic


def notes(slide, body: str):
    slide.notes_slide.notes_text_frame.text = body


# --- templates ---------------------------------------------------------------
def intro_title(prs, episode="Episode 1", category="Category goes here"):
    s = page(prs, CREAM, "Intro title card")
    picture(s, SUN, 830, 160, 260, 260, "Sun logo")
    text(s, 160, 455, 1600, 130, "Wait, I Qualify?!", size=52, color=BARK,
         font=DISPLAY, bold=True, align=PP_ALIGN.CENTER, name="Show title")
    text(s, 260, 580, 1400, 44, "Free Technology Resources for Everyone",
         size=17, color=INK, align=PP_ALIGN.CENTER, name="Tagline")
    rect(s, 840, 644, 240, 4, TEAL, name="Teal rule")
    text(s, 160, 668, 1600, 48, f"{episode}   ·   {category}", size=15,
         color=TEAL_SITE, align=PP_ALIGN.CENTER, name="Episode line")
    rect(s, 0, 1016, W, 64, TEAL_BOX, name="Footer bar")
    text(s, 0, 1016, W, 64, "techempower.org", size=15, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, name="Footer URL")
    notes(s, "Intro title card. Edit the episode line. Holds under the jingle "
             "after the sun-logo pre-roll (see show/PIPELINE.md, stage 8).")
    return s


def lower_third(prs, name_, role, wide=False):
    s = page(prs, None, f"Lower third — {name_}")
    w = 620 if wide else 480
    rect(s, 60, 920, w, 72, TEAL_BOX, name="Lower third box", radius=4)
    rect(s, 60, 920, 6, 72, AMBER, name="Amber accent")
    text(s, 82, 923, w - 40, 40, name_, size=18, color=WHITE, font=DISPLAY,
         bold=True, name="Name")
    text(s, 82, 961, w - 40, 28, role, size=9.5, color=OFFWHITE, name="Role")
    notes(s, "Lower third overlay. Export this page as PNG with 'Transparent "
             "background' ticked so it sits over the interview footage.")
    return s


ROW_PLAIN, ROW_PAIR = 48, 70


def program_card(prs, title, lines, *, cue=None, note=None, label="Program card"):
    s = page(prs, None, f"{label} — {title}")
    body = sum(ROW_PAIR if isinstance(l, tuple) else ROW_PLAIN for l in lines)
    panel_h = 118 + body + 22
    top = 1000 - panel_h
    x, pw = 300, 1320
    rect(s, x, top, pw, panel_h, BROWN, radius=14, name="Card panel")
    rect(s, x, top, 8, panel_h, AMBER, name="Amber accent")
    text(s, x + 44, top + 26, pw - 88, 48, title, size=22, color=WHITE,
         font=DISPLAY, bold=True, name="Program name")
    rect(s, x + 44, top + 86, pw - 88, 2, AMBER, name="Rule")
    y = top + 106
    for i, line in enumerate(lines, 1):
        dot(s, x + 46, y + 11, 12, TEAL, name=f"Bullet {i}")
        if isinstance(line, tuple):
            lab, sub = line
            text(s, x + 72, y, pw - 130, 34, lab, size=15, color=AMBER, bold=True,
                 name=f"Line {i} label")
            text(s, x + 72, y + 32, pw - 130, 30, sub, size=11.5, color=PALE,
                 name=f"Line {i} detail")
            y += ROW_PAIR
        else:
            text(s, x + 72, y, pw - 130, 40, line, size=15, color=PALE,
                 name=f"Line {i}")
            y += ROW_PLAIN
    n = ["Program card overlay. Export as PNG with 'Transparent background' "
         "ticked. Numbers/URLs are from the script cue below, which carries "
         "every fact-check fix; if the audio disagrees, "
         "show/epN/fact-check-full-report.json wins -- flag it to JP."]
    if cue:
        n.append("\nSource cue:\n" + cue)
    if note:
        n.append("\nNote: " + note)
    notes(s, "\n".join(n))
    return s


def url_strap(prs, url="website.example.org"):
    s = page(prs, None, "URL strap (screencast)")
    rect(s, 1200, 956, 660, 68, TEAL_BOX, radius=34, name="Strap")
    dot(s, 1232, 982, 16, AMBER, name="Amber dot")
    text(s, 1262, 956, 570, 68, url, size=16, color=WHITE, bold=True,
         anchor=MSO_ANCHOR.MIDDLE, name="URL")
    notes(s, "URL strap for [POST: SCREENCAST ...] segments -- sits bottom-right "
             "over the screen capture. Transparent-background PNG export.")
    return s


def book_card(prs, kicker="BOOK MENTIONED", title=("Book Title", "Second Line"),
              author="by Author Name", tag="Free at your local library"):
    s = page(prs, None, "Book / resource mentioned")
    rect(s, 1320, 660, 520, 280, CARD, radius=10, edge=TEAL, name="Card")
    text(s, 1345, 674, 470, 26, kicker, size=8.5, color=TEAL, bold=True,
         name="Kicker")
    rect(s, 1345, 707, 470, 2, RULE, name="Rule")
    text(s, 1345, 720, 470, 100, list(title), size=20, color=BROWN, font=DISPLAY,
         bold=True, name="Title")
    text(s, 1345, 826, 470, 34, author, size=12, color=INK, name="Author")
    text(s, 1345, 868, 470, 30, tag, size=9.5, color=TEAL, name="Tag")
    notes(s, "Cream card, lower right. Used in Ep1 for the book mention; reuse "
             "for any single resource callout. Transparent-background export.")
    return s


def find_help(prs):
    s = page(prs, None, "Find Help Today")
    rect(s, 640, 680, 640, 300, BROWN, radius=12, name="Panel")
    text(s, 640, 702, 640, 48, "Find Help Today", size=19, color=WHITE,
         font=DISPLAY, bold=True, align=PP_ALIGN.CENTER, name="Heading")
    rect(s, 700, 762, 520, 2, AMBER, name="Rule")
    items = [("techempower.org", "Free guides & resources"),
             ("Dial 2-1-1", "Local assistance near you"),
             ("findhelp.org", "Search programs by ZIP code")]
    y = 782
    for i, (lab, sub) in enumerate(items, 1):
        dot(s, 680, y + 9, 12, TEAL, name=f"Bullet {i}")
        text(s, 706, y, 500, 32, lab, size=13.5, color=AMBER, bold=True,
             name=f"Item {i}")
        text(s, 706, y + 30, 500, 26, sub, size=9, color=PALE,
             name=f"Item {i} detail")
        y += 64
    notes(s, "Closing resources card -- every episode (series invariant: 211, "
             "findhelp.org, techempower.org). Transparent-background export.")
    return s


def end_card(prs, lines=("techempower.org", "Dial 2-1-1   |   findhelp.org")):
    s = page(prs, DARK, "Outro end card")
    picture(s, SUN, 860, 240, 200, 200, "Sun logo")
    text(s, 160, 462, 1600, 96, "Wait, I Qualify?!", size=34, color=PALE,
         font=DISPLAY, bold=True, align=PP_ALIGN.CENTER, name="Show title")
    lines = list(lines)
    text(s, 160, 556, 1600, 50, lines[0], size=18, color=AMBER, bold=True,
         align=PP_ALIGN.CENTER, name="Primary URL")
    rect(s, 860, 618, 200, 3, TEAL, name="Teal rule")
    y = 640
    for i, line in enumerate(lines[1:], 2):
        text(s, 160, y, 1600, 44, line, size=13.5, color=PALE,
             align=PP_ALIGN.CENTER, name=f"Line {i}")
        y += 48
    text(s, 160, 972, 1600, 36,
         "Free technology resources for low-income individuals and families",
         size=9.5, color=PALE, align=PP_ALIGN.CENTER, name="Tagline")
    notes(s, "Outro end card -- plays under the closing jingle. Series "
             "invariant: 211 + findhelp.org + techempower.org always on it.")
    return s


def chart_frame(prs, title="Chart title goes here",
                source="Source: name the primary source here"):
    s = page(prs, CREAM, "Chart frame")
    text(s, 100, 56, 1600, 72, title, size=21, color=BARK, font=DISPLAY,
         bold=True, name="Chart title")
    rect(s, 100, 136, 120, 4, TEAL, name="Teal rule")
    picture(s, SUN, 1760, 48, 80, 80, "Sun logo")
    rect(s, 100, 168, 1720, 800, PLACEHOLDER, edge=RULE, dashed=True,
         name="Chart image placeholder")
    text(s, 100, 168, 1720, 800, "Drop the chart image here  (1720 × 800)",
         size=12, color=MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
         name="Placeholder label")
    text(s, 100, 992, 1720, 36, source, size=9.5, color=MUTED, name="Source credit")
    notes(s, "Full-frame chart card. Replace the placeholder with the chart "
             "image and credit the source (show/PIPELINE.md series assets).")
    return s


def thumbnail(prs, ep_no, title, hook):
    s = page(prs, CREAM, f"Thumbnail — Ep {ep_no}")
    rect(s, 1180, 0, 740, H, PLACEHOLDER, name="Photo placeholder")
    text(s, 1180, 0, 740, H, ["Host still / photo", "fills this side, edge to edge"],
         size=12, color=MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
         name="Photo label")
    text(s, 100, 110, 1040, 44,
         f"WAIT, I QUALIFY?!   ·   EPISODE {ep_no}", size=13, color=TEAL_SITE,
         bold=True, name="Kicker")
    text(s, 100, 190, 1040, 430, list(hook), size=50, color=BARK, font=DISPLAY,
         bold=True, spacing=1.0, name="Hook")
    text(s, 100, 650, 1040, 60, title, size=17, color=INK, name="Episode title")
    rect(s, 100, 880, 520, 72, TEAL_BOX, radius=36, name="URL pill")
    text(s, 100, 880, 520, 72, "techempower.org", size=16, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, name="URL")
    picture(s, SUN, 980, 870, 130, 130, "Sun logo")
    notes(s, "YouTube thumbnail (export 1280x720). Big hook, one still of the "
             "hosts on the right. Keep text inside the left 1100px so it "
             "survives the small-size crop.")
    return s


def banner(prs):
    # 2560x1440 canvas expressed in the 1920-wide slide space (x0.75).
    s = page(prs, CREAM, "YouTube channel banner")
    # (banner.src.html also has a 5%-opacity ghost sun bleeding off the left
    #  edge; add it in Canva if wanted -- PPTX can't carry image opacity.)
    # safe zone 1546x423 centered -> 1160x317
    sx, sy = (W - 1160) // 2, (H - 317) // 2
    picture(s, SUN, sx, sy + 19, 279, 279, "Sun logo")
    tx = sx + 279 + 58
    text(s, tx, sy - 6, 820, 110, "TechEMPOWER", size=52, color=BARK,
         font=DISPLAY, bold=True, name="Wordmark")
    text(s, tx, sy + 108, 820, 40, "Free technology resources for everyone",
         size=17.5, color=RGBColor(0x5C, 0x42, 0x36), name="Tagline")
    rect(s, tx, sy + 168, 99, 6, TEAL_SITE, radius=3, name="Teal rule")
    text(s, tx, sy + 196, 900, 40,
         "▶  Now streaming: Wait, I Qualify?!   •   techempower.org", size=14,
         color=RGBColor(0x0B, 0x5E, 0x5A), bold=True, name="Feature line")
    notes(s, "Channel banner. Rebuilt from show/assets/youtube/banner.src.html. "
             "Export 2560x1440; everything important sits in YouTube's "
             "1546x423 safe zone. Optional: a big sun at ~5% opacity bleeding "
             "off the left edge, as in the current banner.")
    return s


def avatar(prs):
    s = page(prs, CREAM, "YouTube channel avatar")
    picture(s, SUN, 230, 230, 620, 620, "Sun logo")  # 1080 square space
    notes(s, "Square channel avatar; export 800x800 PNG.")
    return s


# --- decks -------------------------------------------------------------------
def slug(t: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", t.lower()).strip("-")


def build_kit() -> Path:
    prs = new_prs()
    intro_title(prs)
    lower_third(prs, "Jeff Hein", "TechEMPOWER Founder")
    lower_third(prs, "Shawna Hein", "Co-host")
    lower_third(prs, "Guest Name", "Title, Organization", wide=True)
    program_card(prs, "Program name — phone + website",
                 ["Apply: website.example.org", "1-800-555-0100",
                  "Or dial 2-1-1"], label="Program card (simple)")
    program_card(prs, "Program name — address + hours",
                 ["123 Main St, Grass Valley", "530-555-0100",
                  "Mon · Wed · Fri, 10 to 1", "Plus Saturdays"],
                 label="Program card (address + hours)")
    program_card(prs, "Several providers on one card",
                 [("Provider one", "$14.95 a month"),
                  ("Provider two", "$30 or less a month"),
                  ("Provider three", "530-555-0100  ·  book the day before"),
                  "Check your address with each provider"],
                 label="Program card (multi-provider)")
    url_strap(prs)
    book_card(prs)
    find_help(prs)
    end_card(prs)
    chart_frame(prs)
    p = OUT / "wiq-series-template-kit.pptx"
    prs.save(p)
    return p


def build_episode(n: int) -> Path:
    ep = EPISODES[n]
    prs = new_prs()
    intro_title(prs, f"Episode {n}", ep["title"])
    lower_third(prs, "Jeff Hein", "TechEMPOWER Founder")
    lower_third(prs, "Shawna Hein", "Co-host")
    for gname, grole in ep.get("guests", []):
        lower_third(prs, gname, grole, wide=True)
    for c in ep["cards"]:
        program_card(prs, c["title"], c["lines"], cue=c.get("cue"),
                     note=c.get("note"))
    end_card(prs, ep["end_lines"])
    status = f"-{ep['status'].lower()}" if ep.get("status") else ""
    p = OUT / f"wiq-ep{n}-{slug(ep['title'])}{status}.pptx"
    prs.save(p)
    return p


def build_thumbnails() -> Path:
    prs = new_prs()
    for n, ep in sorted({**THUMB_ONLY, **EPISODES}.items()):
        thumbnail(prs, n, ep["title"], ep["thumb_hook"])
    p = OUT / "wiq-youtube-thumbnails.pptx"
    prs.save(p)
    return p


def build_channel_art() -> list[Path]:
    b = new_prs(); banner(b); pb = OUT / "techempower-youtube-banner.pptx"; b.save(pb)
    a = new_prs(square=True); avatar(a); pa = OUT / "techempower-youtube-avatar.pptx"; a.save(pa)
    return [pb, pa]


def main():
    OUT.mkdir(exist_ok=True)
    written = [build_kit(), build_thumbnails(), *build_channel_art()]
    written += [build_episode(n) for n in sorted(EPISODES)]
    for p in written:
        print(f"{p.stat().st_size:>8,d}  {p.relative_to(HERE.parent.parent.parent)}")


if __name__ == "__main__":
    main()
